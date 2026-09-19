"""Разбор присланных файлов.

Три формата — три библиотеки: общего разборщика, который умел бы все, не
существует. Здесь же и единственное место, где эти библиотеки вообще
упоминаются: ядру про них знать нечего.

Всё, что чужая библиотека может бросить на битом файле, приводится к нашим
двум исключениям. Разнообразие способов, какими ломается чужой формат,
бесконечно, а для человека случаев ровно два: «не читается» и «текста нет».
"""

from __future__ import annotations

import io
from collections.abc import Iterator

import pypdf
from docx import Document as DocxDocument
from pptx import Presentation

from app.core.documents import DocumentFormat
from app.core.models import Document
from app.ports.documents import DocumentEmptyError, DocumentUnreadableError


class LocalDocumentReader:
    """Реализация порта DocumentReader. Работает в памяти, файлов не пишет."""

    def read(
        self, document: Document, document_format: DocumentFormat, *, limit: int
    ) -> str:
        if limit <= 0:
            raise ValueError("предел длины текста должен быть положительным")

        source = io.BytesIO(document.data)
        try:
            pieces = _readers[document_format](source)
            text = _joined(pieces, limit)
        except (DocumentEmptyError, DocumentUnreadableError):
            raise
        except Exception as error:
            # Чужой формат ломается чем угодно — от KeyError внутри разбора
            # архива до собственных исключений библиотеки. Для человека это
            # один случай, и звучит он «файл не читается».
            raise DocumentUnreadableError(
                f"файл не открылся: {type(error).__name__}"
            ) from error

        if not text:
            raise DocumentEmptyError("в файле нет текста")
        return text


def _joined(pieces: Iterator[str], limit: int) -> str:
    """Склеивает куски, пока не упрётся в предел.

    Считает по ходу дела и обрывается на пределе: у большого PDF полный текст
    может не поместиться в память там, где урезанный помещается с запасом.
    """
    collected: list[str] = []
    length = 0
    for piece in pieces:
        cleaned = piece.strip()
        if not cleaned:
            continue
        collected.append(cleaned)
        length += len(cleaned) + 1
        if length >= limit:
            break
    return "\n".join(collected)[:limit].strip()


def _from_docx(source: io.BytesIO) -> Iterator[str]:
    """Абзацы и ячейки таблиц.

    Таблицы берутся наравне с текстом: в отчётах и методичках половина цифр
    живёт именно в них, и доклад без них вышел бы про другое.
    """
    document = DocxDocument(source)
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            yield " | ".join(cell.text.strip() for cell in row.cells)


def _from_pptx(source: io.BytesIO) -> Iterator[str]:
    """Текст фигур со всех слайдов, по порядку."""
    for slide in Presentation(source).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                yield shape.text_frame.text


def _from_pdf(source: io.BytesIO) -> Iterator[str]:
    """Текстовый слой страниц.

    У PDF со сканами его нет вовсе: внутри картинки страниц. Мы вернём
    пустоту, а вызывающий скажет человеку, что распознавания у нас нет.
    """
    for page in pypdf.PdfReader(source).pages:
        yield page.extract_text() or ""


_readers = {
    DocumentFormat.DOCX: _from_docx,
    DocumentFormat.PPTX: _from_pptx,
    DocumentFormat.PDF: _from_pdf,
}
