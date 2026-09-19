"""Сборка готового документа из текста модели.

Два формата собираются из одного и того же разбора текста, а не двумя разными
разборами: иначе Word и PDF однажды разошлись бы содержанием, и человек
получил бы два разных доклада под одним именем.

Разметку понимаем самую скромную: строка с решётки — заголовок, строка с тире
или звёздочки — пункт списка, всё прочее — абзац. Больше от модели и не
просим: полноценный markdown тянет за собой таблицы, ссылки и вложенность, а
с ними — разбор, который ломается на первом же кривом символе.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path

from docx import Document as DocxDocument
from fpdf import FPDF

from app.core.documents import DocumentFormat
from app.core.models import Document

#: Шрифт лежит в репозитории, а не берётся из системы: сборка идёт через
#: Nixpacks, и системных шрифтов в образе может не оказаться вовсе — тогда
#: кириллица в PDF превращается в пустые квадраты. См. assets/fonts/README.md.
DEFAULT_FONT = Path("assets/fonts/DejaVuSans.ttf")

#: Имя, под которым шрифт живёт внутри собираемого PDF.
_FONT = "body"

_MIME = {
    DocumentFormat.DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    DocumentFormat.PDF: "application/pdf",
}


class _Kind(StrEnum):
    """Что это за строка разметки."""

    HEADING = "heading"
    BULLET = "bullet"
    PARAGRAPH = "paragraph"


@dataclass(frozen=True, slots=True)
class _Line:
    kind: _Kind
    text: str
    #: Глубина заголовка: одна решётка — первая, две — вторая. У прочих ноль.
    level: int = 0


class LocalDocumentWriter:
    """Реализация порта DocumentWriter. Собирает в памяти, файлов не пишет."""

    def __init__(self, font: Path = DEFAULT_FONT) -> None:
        self._font = font

    def build(
        self, title: str, body: str, *, document_format: DocumentFormat
    ) -> Document:
        if document_format is DocumentFormat.PPTX:
            # Презентацию мы пока не собираем, и промолчать тут нельзя:
            # человек получил бы пустой файл вместо доклада.
            raise ValueError("сборка pptx не поддерживается")

        lines = _parse(body)
        data = (
            _docx(title, lines)
            if document_format is DocumentFormat.DOCX
            else _pdf(title, lines, self._font)
        )
        return Document(
            data=data,
            filename=f"{_safe_name(title)}.{document_format.value}",
            mime_type=_MIME[document_format],
        )


def _parse(body: str) -> list[_Line]:
    """Разбирает текст модели в строки разметки."""
    lines: list[_Line] = []
    for raw in body.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("#"):
            hashes = len(line) - len(line.lstrip("#"))
            text = line.lstrip("#").strip()
            if text:
                lines.append(_Line(_Kind.HEADING, text, level=min(hashes, 3)))
            continue
        if line[0] in "-*•" and len(line) > 1:
            lines.append(_Line(_Kind.BULLET, line[1:].strip()))
            continue
        lines.append(_Line(_Kind.PARAGRAPH, line))
    return lines


def _docx(title: str, lines: list[_Line]) -> bytes:
    document = DocxDocument()
    document.add_heading(title, 0)
    for line in lines:
        if line.kind is _Kind.HEADING:
            document.add_heading(line.text, line.level)
        elif line.kind is _Kind.BULLET:
            document.add_paragraph(line.text, style="List Bullet")
        else:
            document.add_paragraph(line.text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(title: str, lines: list[_Line], font: Path) -> bytes:
    pdf = FPDF()
    # Шрифт подключается явно и один: встроенные шрифты fpdf2 кириллицу не
    # знают, и доклад вышел бы страницей вопросительных знаков.
    pdf.add_font(_FONT, "", str(font))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    pdf.set_font(_FONT, size=18)
    pdf.multi_cell(0, 10, title)
    pdf.ln(4)

    for line in lines:
        if line.kind is _Kind.HEADING:
            pdf.ln(3)
            pdf.set_font(_FONT, size=15 if line.level <= 1 else 13)
            pdf.multi_cell(0, 8, line.text)
            pdf.ln(1)
            continue
        pdf.set_font(_FONT, size=11)
        pdf.multi_cell(
            0, 6, f"• {line.text}" if line.kind is _Kind.BULLET else line.text
        )
        pdf.ln(1)

    return bytes(pdf.output())


def _safe_name(title: str) -> str:
    """Имя файла из заголовка.

    Мессенджеры и файловые системы спотыкаются о слэши и двоеточия, а человеку
    важно узнать файл в списке загрузок. Поэтому режем опасное, а не
    придумываем имя заново.
    """
    cleaned = "".join(
        " " if letter in '\\/:*?"<>|\n\r\t' else letter for letter in title
    )
    trimmed = " ".join(cleaned.split())[:60].strip()
    return trimmed or "Документ"
