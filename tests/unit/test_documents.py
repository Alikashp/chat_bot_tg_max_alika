"""Чтение присланных файлов и сборка готовых.

Проверяется не «функция вызвалась», а круг целиком: собрали файл настоящей
библиотекой, прочитали его обратно другой и убедились, что дошло то же самое.
Кириллица здесь не придирка — на ней разваливается и PDF без подключённого
шрифта, и любая попытка обойтись латиницей в тестах.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader

from app.adapters.documents.reader import LocalDocumentReader
from app.adapters.documents.writer import LocalDocumentWriter
from app.core.documents import (
    DocumentFormat,
    DocumentProblem,
    check_document,
    detect_format,
)
from app.core.models import Document
from app.ports.documents import DocumentEmptyError, DocumentUnreadableError

RU = "Влияние климата на урожай: выводы и цифры"


def _docx(*paragraphs: str, table: list[list[str]] | None = None) -> Document:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        added = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(added.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return Document(
        data=buffer.getvalue(), filename="исходник.docx", mime_type="application/x"
    )


def _pptx(*titles: str) -> Document:
    presentation = Presentation()
    for title in titles:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
    buffer = io.BytesIO()
    presentation.save(buffer)
    return Document(
        data=buffer.getvalue(), filename="слайды.pptx", mime_type="application/x"
    )


def _pdf(body: str) -> Document:
    built = LocalDocumentWriter().build(body, "", document_format=DocumentFormat.PDF)
    return Document(data=built.data, filename="файл.pdf", mime_type="application/pdf")


# --- Опознание формата ---------------------------------------------------


def test_a_docx_is_recognised() -> None:
    assert detect_format(_docx("текст")) is DocumentFormat.DOCX


def test_a_pptx_is_told_apart_from_a_docx() -> None:
    """Оба внутри zip: по сигнатуре они неразличимы, различает расширение."""
    assert detect_format(_pptx("слайд")) is DocumentFormat.PPTX


def test_a_pdf_is_recognised() -> None:
    assert detect_format(_pdf("текст")) is DocumentFormat.PDF


def test_a_renamed_file_is_refused() -> None:
    """Под именем «доклад.pdf» может приехать что угодно — сигнатура не сойдётся."""
    disguised = Document(
        data=_docx("текст").data, filename="доклад.pdf", mime_type="application/pdf"
    )

    assert detect_format(disguised) is None


def test_an_unknown_extension_is_refused() -> None:
    assert (
        detect_format(Document(b"PK\x03\x04", "архив.zip", "application/zip")) is None
    )


def test_a_file_without_an_extension_is_refused() -> None:
    assert detect_format(Document(b"%PDF-1.7", "документ", "application/pdf")) is None


def test_a_file_that_is_too_big_is_refused_before_its_format() -> None:
    """Человеку важнее узнать про размер: формат он менять не станет."""
    big = Document(data=b"x" * 100, filename="что-то.txt", mime_type="text/plain")

    assert check_document(big, max_bytes=10).problem is DocumentProblem.TOO_BIG


def test_a_good_file_passes_with_its_format() -> None:
    check = check_document(_docx("текст"), max_bytes=10_000_000)

    assert check.ok
    assert check.document_format is DocumentFormat.DOCX


# --- Чтение --------------------------------------------------------------


def test_text_comes_out_of_a_docx() -> None:
    text = LocalDocumentReader().read(
        _docx(RU, "второй абзац"), DocumentFormat.DOCX, limit=10_000
    )

    assert RU in text
    assert "второй абзац" in text


def test_a_table_is_read_too() -> None:
    """В отчётах половина цифр живёт в таблицах — доклад без них про другое."""
    source = _docx("шапка", table=[["Год", "Урожай"], ["2026", "41 центнер"]])

    text = LocalDocumentReader().read(source, DocumentFormat.DOCX, limit=10_000)

    assert "41 центнер" in text


def test_text_comes_out_of_a_pptx() -> None:
    text = LocalDocumentReader().read(
        _pptx(RU, "второй слайд"), DocumentFormat.PPTX, limit=10_000
    )

    assert RU in text
    assert "второй слайд" in text


def test_text_comes_out_of_a_pdf() -> None:
    text = LocalDocumentReader().read(_pdf(RU), DocumentFormat.PDF, limit=10_000)

    assert "климата" in text


def test_the_text_is_cut_at_the_limit() -> None:
    """Предел бережёт и память, и деньги: длинный файл — дорогой запрос."""
    text = LocalDocumentReader().read(
        _docx("а" * 500, "б" * 500), DocumentFormat.DOCX, limit=100
    )

    assert len(text) <= 100


def test_a_broken_file_is_named_as_unreadable() -> None:
    """Битый архив, пароль, чужой формат — для человека это один случай."""
    broken = Document(b"PK\x03\x04" + "мусор".encode(), "битый.docx", "application/x")

    with pytest.raises(DocumentUnreadableError):
        LocalDocumentReader().read(broken, DocumentFormat.DOCX, limit=10_000)


def test_a_file_without_a_text_layer_is_named_as_empty() -> None:
    """Ровно случай PDF-скана: страницы есть, текста нет.

    Отправлять это провайдеру нельзя: он сочинит доклад на вольную тему, а
    человек заплатит за него лимитом.
    """
    with pytest.raises(DocumentEmptyError):
        LocalDocumentReader().read(_docx("   "), DocumentFormat.DOCX, limit=10_000)


def test_a_limit_of_zero_is_a_mistake_not_an_empty_answer() -> None:
    with pytest.raises(ValueError):
        LocalDocumentReader().read(_docx("текст"), DocumentFormat.DOCX, limit=0)


# --- Сборка --------------------------------------------------------------

BODY = "# Выводы\n\nПервый абзац.\n\n- первый пункт\n- второй пункт"


def test_a_built_docx_reads_back() -> None:
    built = LocalDocumentWriter().build(RU, BODY, document_format=DocumentFormat.DOCX)

    text = LocalDocumentReader().read(built, DocumentFormat.DOCX, limit=10_000)
    assert RU in text
    assert "Выводы" in text
    assert "первый пункт" in text


def test_a_built_pdf_keeps_cyrillic() -> None:
    """Без подключённого шрифта здесь были бы пустые квадраты."""
    built = LocalDocumentWriter().build(RU, BODY, document_format=DocumentFormat.PDF)

    text = "".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(built.data)).pages
    )
    assert "климата" in text
    assert "первый пункт" in text


def test_both_formats_carry_the_same_words() -> None:
    """Иначе человек получил бы два разных доклада под одним именем."""
    writer = LocalDocumentWriter()

    as_docx = writer.build(RU, BODY, document_format=DocumentFormat.DOCX)
    as_pdf = writer.build(RU, BODY, document_format=DocumentFormat.PDF)

    from_docx = LocalDocumentReader().read(as_docx, DocumentFormat.DOCX, limit=10_000)
    from_pdf = "".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(as_pdf.data)).pages
    )
    for word in ("Выводы", "Первый абзац.", "второй пункт"):
        assert word in from_docx
        assert word in from_pdf


def test_a_heading_line_does_not_keep_its_hashes() -> None:
    built = LocalDocumentWriter().build(
        RU, "## Раздел", document_format=DocumentFormat.DOCX
    )

    text = LocalDocumentReader().read(built, DocumentFormat.DOCX, limit=10_000)
    assert "Раздел" in text
    assert "#" not in text


def test_the_file_name_comes_from_the_title() -> None:
    built = LocalDocumentWriter().build(RU, BODY, document_format=DocumentFormat.DOCX)

    assert built.filename.endswith(".docx")
    assert "климата" in built.filename


def test_a_title_with_slashes_still_makes_a_usable_name() -> None:
    """Мессенджеры и файловые системы спотыкаются о слэши и двоеточия."""
    built = LocalDocumentWriter().build(
        "Отчёт 1/2: итоги", BODY, document_format=DocumentFormat.PDF
    )

    assert "/" not in built.filename
    assert ":" not in built.filename
    assert built.filename.endswith(".pdf")


def test_a_built_docx_is_a_real_office_file() -> None:
    """Word открывает только настоящий OOXML, а не что-то похожее."""
    built = LocalDocumentWriter().build(RU, BODY, document_format=DocumentFormat.DOCX)

    with zipfile.ZipFile(io.BytesIO(built.data)) as archive:
        assert "word/document.xml" in archive.namelist()


def test_building_a_pptx_is_refused_rather_than_returning_an_empty_file() -> None:
    """Промолчать нельзя: человек получил бы пустой файл вместо доклада."""
    with pytest.raises(ValueError):
        LocalDocumentWriter().build(RU, BODY, document_format=DocumentFormat.PPTX)
