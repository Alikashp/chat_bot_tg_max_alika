"""Раздел «Документы»: от выбора действия до двух готовых файлов.

Главное, что здесь доказывается, — ключевой инвариант проекта: за упавший
разбор лимит не списывается. Проверяется он не на словах, а по остатку до и
после каждого отказа: битый файл, скан, сбой провайдера.

Второе — что действие остаётся действием. Нажал «Реферат» — в инструкции
провайдеру должен оказаться реферат, а в отданных файлах — два формата, а не
один.
"""

from __future__ import annotations

import io
from dataclasses import replace

import pytest
from docx import Document as DocxDocument

from app.adapters.storage.memory import InMemoryStorage
from app.core.limits import LimitKind
from app.core.models import Document
from app.core.pending import await_document, parse_await_document
from app.core.scenarios import documents, spending
from app.core.scenarios.deps import Deps, Session
from config.documents import DOCUMENT_ACTIONS
from tests.fakes import FakeLLM, FakeMessenger

REPORT = DOCUMENT_ACTIONS["report"]
ABSTRACT = DOCUMENT_ACTIONS["abstract"]


def _docx(text: str = "Текст исходного документа про урожай 2026 года.") -> Document:
    document = DocxDocument()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return Document(
        data=buffer.getvalue(), filename="исходник.docx", mime_type="application/x"
    )


async def _left(deps: Deps, session: Session) -> int:
    allowance = await spending.current_allowance(deps, session, LimitKind.DOCUMENTS)
    return allowance.total_left


async def _with_documents(
    storage: InMemoryStorage, session: Session, count: int
) -> Session:
    await storage.add_bonus(session.user.id, documents=count)
    fresh = await storage.get_user_by_id(session.user.id)
    assert fresh is not None
    return replace(session, user=fresh)


# --- Выбор действия ------------------------------------------------------


async def test_the_menu_offers_every_action_from_the_registry(
    deps: Deps, session: Session, messenger: FakeMessenger
) -> None:
    """Новое действие — запись в реестре, и в меню оно приходит само."""
    await documents.show_menu(deps, session)

    keyboard = messenger.last_text.keyboard
    assert keyboard is not None
    labels = [button.text for row in keyboard.rows for button in row]
    assert labels == [action.button for action in DOCUMENT_ACTIONS.values()]


async def test_choosing_an_action_asks_for_a_file(
    deps: Deps, session: Session, storage: InMemoryStorage, messenger: FakeMessenger
) -> None:
    ready = await _with_documents(storage, session, 1)

    await documents.choose(deps, ready, REPORT)

    assert messenger.last_text.text == REPORT.invitation
    fresh = await storage.get_user_by_id(ready.user.id)
    assert fresh is not None
    assert parse_await_document(fresh.pending) == REPORT.id


async def test_an_exhausted_person_is_told_before_sending_a_file(
    deps: Deps, session: Session, storage: InMemoryStorage, messenger: FakeMessenger
) -> None:
    """Иначе человек выберет действие, дождётся загрузки и узнает только тогда."""
    await documents.choose(deps, session, REPORT)

    fresh = await storage.get_user_by_id(session.user.id)
    assert fresh is not None
    assert fresh.pending is None
    assert messenger.last_text.text != REPORT.invitation


# --- Разбор --------------------------------------------------------------


async def test_a_file_comes_back_as_two_files(
    deps: Deps, session: Session, storage: InMemoryStorage, messenger: FakeMessenger
) -> None:
    """Word и PDF сразу: выбор формата ничего не стоит, а экран стоит нажатия."""
    ready = await _with_documents(storage, session, 1)

    await documents.apply(deps, ready, REPORT, _docx())

    names = [document.filename for document in messenger.documents_sent]
    assert len(names) == 2
    assert any(name.endswith(".docx") for name in names)
    assert any(name.endswith(".pdf") for name in names)


async def test_the_chosen_action_reaches_the_provider(
    deps: Deps, session: Session, storage: InMemoryStorage, llm: FakeLLM
) -> None:
    """Нажал «Реферат» — обязан получить реферат, а не доклад."""
    ready = await _with_documents(storage, session, 1)

    await documents.apply(deps, ready, ABSTRACT, _docx())

    turns, _ = llm.calls[0]
    assert ABSTRACT.instruction in turns[0].content


async def test_the_text_of_the_file_reaches_the_provider(
    deps: Deps, session: Session, storage: InMemoryStorage, llm: FakeLLM
) -> None:
    ready = await _with_documents(storage, session, 1)

    await documents.apply(deps, ready, REPORT, _docx("урожай 41 центнер"))

    turns, _ = llm.calls[0]
    assert "41 центнер" in turns[0].content


async def test_a_finished_analysis_costs_one(
    deps: Deps, session: Session, storage: InMemoryStorage
) -> None:
    ready = await _with_documents(storage, session, 3)
    before = await _left(deps, ready)

    await documents.apply(deps, ready, REPORT, _docx())

    assert await _left(deps, ready) == before - 1


# --- За упавшее не платят ------------------------------------------------


async def test_a_broken_file_costs_nothing(
    deps: Deps, session: Session, storage: InMemoryStorage, llm: FakeLLM
) -> None:
    ready = await _with_documents(storage, session, 2)
    before = await _left(deps, ready)
    broken = Document(b"PK\x03\x04" + b"mycop", "битый.docx", "application/x")

    await documents.apply(deps, ready, REPORT, broken)

    assert await _left(deps, ready) == before
    assert llm.calls == []


async def test_a_scan_costs_nothing_and_says_so(
    deps: Deps,
    session: Session,
    storage: InMemoryStorage,
    messenger: FakeMessenger,
    llm: FakeLLM,
) -> None:
    """У скана нет текстового слоя: провайдер сочинил бы доклад на вольную тему."""
    ready = await _with_documents(storage, session, 2)
    before = await _left(deps, ready)

    await documents.apply(deps, ready, REPORT, _docx("   "))

    assert await _left(deps, ready) == before
    assert llm.calls == []
    assert "скан" in messenger.last_text.text


async def test_an_unsupported_file_costs_nothing(
    deps: Deps, session: Session, storage: InMemoryStorage, llm: FakeLLM
) -> None:
    ready = await _with_documents(storage, session, 2)
    before = await _left(deps, ready)
    alien = Document(b"just text", "заметка.txt", "text/plain")

    await documents.apply(deps, ready, REPORT, alien)

    assert await _left(deps, ready) == before
    assert llm.calls == []


async def test_a_provider_failure_costs_nothing(
    deps: Deps, session: Session, storage: InMemoryStorage, llm: FakeLLM
) -> None:
    """Ключевой инвариант: лимит не списывается за упавший запрос."""
    ready = await _with_documents(storage, session, 2)
    before = await _left(deps, ready)
    llm.error = RuntimeError("провайдер лёг")

    await documents.apply(deps, ready, REPORT, _docx())

    assert await _left(deps, ready) == before


async def test_a_provider_failure_sends_no_files(
    deps: Deps,
    session: Session,
    storage: InMemoryStorage,
    messenger: FakeMessenger,
    llm: FakeLLM,
) -> None:
    ready = await _with_documents(storage, session, 2)
    llm.error = RuntimeError("провайдер лёг")

    await documents.apply(deps, ready, REPORT, _docx())

    assert messenger.documents_sent == []


async def test_an_exhausted_person_never_reaches_the_provider(
    deps: Deps, session: Session, llm: FakeLLM
) -> None:
    await documents.apply(deps, session, REPORT, _docx())

    assert llm.calls == []


async def test_a_failed_delivery_costs_nothing(
    deps: Deps, session: Session, storage: InMemoryStorage, messenger: FakeMessenger
) -> None:
    """Списание — только по факту доставленного человеку результата.

    Провайдер отработал и деньги за вызов уплачены, но человек файла не
    увидел. Списать здесь значило бы взять с него за то, чего он не получил.
    """
    ready = await _with_documents(storage, session, 2)
    before = await _left(deps, ready)
    messenger.fail_send_document = RuntimeError("мессенджер не принял файл")

    with pytest.raises(RuntimeError):
        await documents.apply(deps, ready, REPORT, _docx())

    assert await _left(deps, ready) == before


# --- Ожидание ------------------------------------------------------------


async def test_the_wait_is_cleared_before_the_provider_is_called(
    deps: Deps, session: Session, storage: InMemoryStorage, llm: FakeLLM
) -> None:
    """Разбор может упасть, а следующий файл не должен приклеиться к прошлому."""
    ready = await _with_documents(storage, session, 2)
    await storage.set_pending(ready.user.id, await_document(REPORT.id))
    llm.error = RuntimeError("провайдер лёг")

    await documents.apply(deps, ready, REPORT, _docx())

    fresh = await storage.get_user_by_id(ready.user.id)
    assert fresh is not None
    assert fresh.pending is None


async def test_a_rejected_file_keeps_the_chosen_action(
    deps: Deps, session: Session, storage: InMemoryStorage
) -> None:
    """Человек уже выбрал, что хотел: неподходящий файл не должен это стирать."""
    ready = await _with_documents(storage, session, 2)
    await storage.set_pending(ready.user.id, await_document(REPORT.id))

    await documents.apply(deps, ready, REPORT, Document(b"x", "z.txt", "text/plain"))

    fresh = await storage.get_user_by_id(ready.user.id)
    assert fresh is not None
    assert parse_await_document(fresh.pending) == REPORT.id
